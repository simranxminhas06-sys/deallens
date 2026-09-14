"""Assembles a slide-deck version of the same analysis generate_report() renders as
Markdown — an actual IC/partner deck (title, verdict, economics, opportunities, risks,
100-day plan) rather than a linear document, since that's the artifact a deal team
hands upward. Pure python-pptx, no PowerPoint installation required, so this deploys
the same way everywhere the rest of the app does.

Visual language mirrors the Streamlit app's own theme (navy/charcoal + bright blue
accent, card-style metrics) rather than python-pptx's default Office look — a thin
accent strip + wordmark + page number on every slide, rounded stat cards instead of
plain bullets for numbers, and table styling that doesn't leak PowerPoint's default
blue theme through our own colors.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from agent.verdict import compute_verdict
from schemas.analysis_models import AnalysisRecord
from tools.financial_calculator import calculate_deal_economics

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
DARK_BG = RGBColor(0x0B, 0x12, 0x20)
CARD_BG = RGBColor(0x14, 0x1B, 0x2E)
CARD_BORDER = RGBColor(0x2A, 0x3B, 0x5C)
ACCENT = RGBColor(0x3B, 0x7A, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xE7, 0xEC, 0xF5)
MUTED = RGBColor(0x8A, 0x96, 0xAC)
ROW_ALT = RGBColor(0x11, 0x17, 0x27)

VERDICT_COLOR = {
    "proceed": RGBColor(0x1F, 0xA9, 0x71),
    "proceed_with_conditions": RGBColor(0xF5, 0xA6, 0x23),
    "further_diligence": RGBColor(0x8B, 0x6B, 0xF2),
    "do_not_proceed": RGBColor(0xE5, 0x48, 0x4D),
}

VERDICT_LABEL = {
    "proceed": "Proceed",
    "proceed_with_conditions": "Proceed with conditions",
    "further_diligence": "Further diligence required",
    "do_not_proceed": "Do not proceed",
}

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.5)
CONTENT_W = SLIDE_W - 2 * MARGIN
CARD_RADIUS = 0.06


def _no_line(shape):
    shape.line.fill.background()
    shape.shadow.inherit = False


def _set_cell_border(cell, color=CARD_BORDER, weight_pt=0.75):
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        existing = tc_pr.find(qn(tag))
        if existing is not None:
            tc_pr.remove(existing)
        ln = OxmlElement(tag)
        ln.set("w", str(int(Pt(weight_pt))))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        fill = OxmlElement("a:solidFill")
        clr = OxmlElement("a:srgbClr")
        clr.set("val", "%02X%02X%02X" % (color[0], color[1], color[2]))
        fill.append(clr)
        ln.append(fill)
        tc_pr.append(ln)


def _blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    _no_line(bg)
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)
    return slide


def _content_slide(prs: Presentation, page_num: int, deal_label: str):
    """A slide with the shared chrome: top accent strip, wordmark, and a footer
    (deal name + page number) — everything after the title slide uses this."""
    slide = _blank_slide(prs)

    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Pt(4))
    strip.fill.solid()
    strip.fill.fore_color.rgb = ACCENT
    _no_line(strip)

    wm = slide.shapes.add_textbox(SLIDE_W - Inches(3) - MARGIN, Inches(0.28), Inches(3), Inches(0.3))
    wp = wm.text_frame.paragraphs[0]
    wp.text = "D E A L L E N S"
    wp.font.size = Pt(10)
    wp.font.bold = True
    wp.font.color.rgb = MUTED
    wp.alignment = PP_ALIGN.RIGHT

    fl = slide.shapes.add_textbox(MARGIN, SLIDE_H - Inches(0.42), Inches(8), Inches(0.3))
    flp = fl.text_frame.paragraphs[0]
    flp.text = deal_label
    flp.font.size = Pt(9)
    flp.font.color.rgb = MUTED

    fr = slide.shapes.add_textbox(SLIDE_W - MARGIN - Inches(1), SLIDE_H - Inches(0.42), Inches(1), Inches(0.3))
    frp = fr.text_frame.paragraphs[0]
    frp.text = f"{page_num:02d}"
    frp.font.size = Pt(9)
    frp.font.color.rgb = MUTED
    frp.alignment = PP_ALIGN.RIGHT

    return slide


def _title(slide, text: str, subtitle: str | None = None):
    box = slide.shapes.add_textbox(MARGIN, Inches(0.32), CONTENT_W, Inches(0.55))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = WHITE

    underline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(0.88), Inches(0.55), Pt(3))
    underline.fill.solid()
    underline.fill.fore_color.rgb = ACCENT
    _no_line(underline)

    if subtitle:
        sub = slide.shapes.add_textbox(MARGIN, Inches(1.0), CONTENT_W, Inches(0.4))
        sp = sub.text_frame.paragraphs[0]
        sub.text_frame.word_wrap = True
        sp.text = subtitle
        sp.font.size = Pt(12)
        sp.font.italic = True
        sp.font.color.rgb = MUTED
        return Inches(1.5)
    return Inches(1.25)


def _body_textbox(slide, top, height=None):
    height = height or (SLIDE_H - top - Inches(0.55))
    box = slide.shapes.add_textbox(MARGIN, top, CONTENT_W, height)
    box.text_frame.word_wrap = True
    return box


def _bullets(slide, lines: list[str], top, font_size=14, height=None, marker_color=ACCENT):
    box = _body_textbox(slide, top, height)
    tf = box.text_frame
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        indented = line.startswith("  ")
        marker = "–" if indented else "▸"
        run_marker = p.add_run()
        run_marker.text = f"{'    ' if indented else ''}{marker}  "
        run_marker.font.size = Pt(font_size)
        run_marker.font.bold = True
        run_marker.font.color.rgb = marker_color
        run_text = p.add_run()
        run_text.text = line.strip()
        run_text.font.size = Pt(font_size)
        run_text.font.color.rgb = LIGHT_GRAY
        p.space_after = Pt(8)
    return box


def _card(slide, x, y, w, h):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.adjustments[0] = CARD_RADIUS
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(1)
    card.shadow.inherit = False
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.14)
    tf.margin_bottom = Inches(0.14)
    return card


def _stat_card(slide, x, y, w, h, label: str, value: str, value_color=None):
    card = _card(slide, x, y, w, h)
    tf = card.text_frame
    p1 = tf.paragraphs[0]
    p1.text = label.upper()
    p1.font.size = Pt(11)
    p1.font.bold = True
    p1.font.color.rgb = MUTED
    p2 = tf.add_paragraph()
    p2.text = value
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = value_color or WHITE
    p2.space_before = Pt(8)
    return card


def _table(slide, header: list[str], rows: list[list[str]], top, col_widths=None, font_size=10):
    n_rows, n_cols = len(rows) + 1, len(header)
    height = min(Inches(0.4) * n_rows, SLIDE_H - top - Inches(0.55))
    shape = slide.shapes.add_table(n_rows, n_cols, MARGIN, top, CONTENT_W, height)
    table = shape.table
    table.first_row = False
    table.horz_banding = False
    if col_widths:
        total = sum(col_widths)
        for i, w in enumerate(col_widths):
            table.columns[i].width = int(CONTENT_W * (w / total))

    for c, h in enumerate(header):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.margin_left, cell.margin_right = Inches(0.08), Inches(0.08)
        cell.margin_top, cell.margin_bottom = Inches(0.05), Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(font_size)
        p.font.color.rgb = WHITE
        _set_cell_border(cell)

    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r % 2 else ROW_ALT
            cell.margin_left, cell.margin_right = Inches(0.08), Inches(0.08)
            cell.margin_top, cell.margin_bottom = Inches(0.05), Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(font_size)
            p.font.color.rgb = LIGHT_GRAY
            _set_cell_border(cell)
    return shape


def _truncate_rows(rows: list[list[str]], limit: int) -> list[list[str]]:
    if len(rows) <= limit:
        return rows
    return rows[: limit - 1] + [[f"+ {len(rows) - limit + 1} more"] + [""] * (len(rows[0]) - 1)]


def generate_pptx(record: AnalysisRecord) -> bytes:
    t = record.transaction
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    deal_label = f"{t.acquirer_name} / {t.target_name}"
    total_value_creation = sum(o.estimated_value.base for o in record.opportunities)
    page = [1]

    def new_slide():
        page[0] += 1
        return _content_slide(prs, page[0], deal_label)

    # --- Title slide
    slide = _blank_slide(prs)
    accent_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(2.55), Inches(1), Pt(4))
    accent_line.fill.solid()
    accent_line.fill.fore_color.rgb = ACCENT
    _no_line(accent_line)
    accent_line.left = int((SLIDE_W - Inches(1)) / 2)

    box = slide.shapes.add_textbox(MARGIN, Inches(2.8), CONTENT_W, Inches(2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "DealLens Analysis"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = deal_label
    p2.font.size = Pt(22)
    p2.font.color.rgb = ACCENT
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(6)
    p3 = tf.add_paragraph()
    p3.text = f"Generated {record.created_at:%Y-%m-%d %H:%M UTC}"
    p3.font.size = Pt(12)
    p3.font.color.rgb = MUTED
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(14)

    # --- Executive summary
    slide = new_slide()
    top = _title(slide, "Executive Summary")
    card = _card(slide, MARGIN, top, CONTENT_W, SLIDE_H - top - Inches(0.65))
    p = card.text_frame.paragraphs[0]
    p.text = record.executive_summary or "Not yet generated."
    p.font.size = Pt(15)
    p.font.color.rgb = LIGHT_GRAY

    # --- Recommendation / verdict
    if record.review is not None:
        verdict = compute_verdict(record.review, record.agent_assessments, total_value_creation, t)
        slide = new_slide()
        top = _title(
            slide,
            "Recommendation",
            "A rule-based verdict, not a model's opinion — every reason traces to a specific reviewer finding or Red-Team challenge.",
        )
        banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, MARGIN, top, CONTENT_W, Inches(0.65))
        banner.adjustments[0] = 0.12
        banner.fill.solid()
        banner.fill.fore_color.rgb = VERDICT_COLOR[verdict.level.value]
        _no_line(banner)
        bp = banner.text_frame.paragraphs[0]
        bp.text = VERDICT_LABEL[verdict.level.value]
        bp.font.size = Pt(19)
        bp.font.bold = True
        bp.font.color.rgb = DARK_BG
        bp.alignment = PP_ALIGN.CENTER
        banner.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

        _bullets(slide, verdict.reasons, top=top + Inches(0.9), font_size=13)
        if verdict.conditions:
            cond_top = top + Inches(0.9) + Inches(0.35) * len(verdict.reasons) + Inches(0.3)
            cond_top = min(cond_top, SLIDE_H - Inches(2.2))
            _bullets(
                slide,
                ["Conditions to resolve before proceeding:"] + [f"  {c}" for c in verdict.conditions],
                top=cond_top,
                font_size=12,
                marker_color=VERDICT_COLOR[verdict.level.value],
            )

    # --- Deal economics
    if t.deal_value and record.opportunities:
        econ = calculate_deal_economics(total_value_creation, t.deal_value)
        slide = new_slide()
        top = _title(slide, "Deal Economics")
        card_w = (CONTENT_W - Inches(0.6)) / 3
        card_h = Inches(1.6)
        stats = [
            ("Value creation", f"${total_value_creation:,.0f}", ACCENT),
            ("Deal value", f"${t.deal_value:,.0f}", WHITE),
            ("% of deal value", f"{econ['value_creation_pct_of_deal']:.2f}%", ACCENT),
        ]
        for i, (label, value, color) in enumerate(stats):
            x = MARGIN + i * (card_w + Inches(0.3))
            _stat_card(slide, x, top, card_w, card_h, label, value, value_color=color)

    # --- Company profiles
    slide = new_slide()
    top = _title(slide, "Company Profiles")
    col_w = (CONTENT_W - Inches(0.4)) / 2
    left_positions = [MARGIN, MARGIN + col_w + Inches(0.4)]
    card_h = SLIDE_H - top - Inches(0.65)
    for pos, (label, profile) in zip(
        left_positions, (("Acquirer", record.acquirer_profile), ("Target", record.target_profile))
    ):
        if not profile:
            continue
        card = _card(slide, pos, top, col_w, card_h)
        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP

        p0 = tf.paragraphs[0]
        p0.text = label.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT

        p1 = tf.add_paragraph()
        p1.text = profile.company_name
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p1.space_before = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = profile.business_description
        p2.font.size = Pt(12)
        p2.font.color.rgb = LIGHT_GRAY
        p2.space_before = Pt(10)

        stats = []
        if profile.revenue is not None:
            stats.append(f"Revenue: ${profile.revenue:,.0f}")
        if profile.revenue_growth_rate is not None:
            stats.append(f"Growth: {profile.revenue_growth_rate * 100:.1f}%")
        if profile.operating_margin is not None:
            stats.append(f"Operating margin: {profile.operating_margin * 100:.1f}%")
        for s in stats:
            ps = tf.add_paragraph()
            ps.text = s
            ps.font.size = Pt(11)
            ps.font.bold = True
            ps.font.color.rgb = MUTED
            ps.space_before = Pt(6)

    # --- Strategic rationale
    if record.strategic_rationale:
        slide = new_slide()
        top = _title(slide, "Strategic Rationale")
        card = _card(slide, MARGIN, top, CONTENT_W, SLIDE_H - top - Inches(0.65))
        p = card.text_frame.paragraphs[0]
        p.text = record.strategic_rationale.summary
        p.font.size = Pt(15)
        p.font.color.rgb = LIGHT_GRAY

    # --- Financial baseline
    if record.financial_baselines:
        slide = new_slide()
        top = _title(slide, "Financial Baseline")
        rows = [[b.metric, f"{b.value:,.2f} {b.unit}", b.method] for b in record.financial_baselines]
        _table(slide, ["Metric", "Value", "Method"], _truncate_rows(rows, 10), top=top, col_widths=[2, 2, 4])

    # --- Opportunities
    revenue_opps = [o for o in record.opportunities if o.category.value == "revenue_synergy"]
    cost_opps = [o for o in record.opportunities if o.category.value == "cost_synergy"]
    for section_title, opps in (("Revenue Opportunities", revenue_opps), ("Cost-Saving Opportunities", cost_opps)):
        if not opps:
            continue
        slide = new_slide()
        top = _title(slide, section_title)
        rows = [
            [
                o.title,
                f"${o.estimated_value.low:,.0f} / ${o.estimated_value.base:,.0f} / ${o.estimated_value.high:,.0f}",
                o.implementation_difficulty.value,
                o.time_horizon,
            ]
            for o in opps
        ]
        _table(
            slide,
            ["Opportunity", "Low / Base / High", "Difficulty", "Horizon"],
            _truncate_rows(rows, 8),
            top=top,
            col_widths=[3, 3, 1.5, 1.5],
        )

    # --- Risk register
    if record.risks:
        slide = new_slide()
        top = _title(slide, "Risk Register", "Sorted by risk score (likelihood x impact, 1-9), highest first.")
        rows = [
            [r.title, r.category, r.likelihood.value, r.severity.value, str(r.score)]
            for r in sorted(record.risks, key=lambda r: r.score, reverse=True)
        ]
        _table(
            slide,
            ["Risk", "Category", "Likelihood", "Impact", "Score"],
            _truncate_rows(rows, 10),
            top=top,
            col_widths=[3, 2, 1.5, 1.5, 1],
        )

    # --- 100-day plan
    if record.integration_plan:
        slide = new_slide()
        top = _title(slide, "100-Day Integration Plan")
        lines = []
        for phase in ("0-30 days", "31-60 days", "61-100 days"):
            actions = [a for a in record.integration_plan.actions if a.phase == phase]
            if not actions:
                continue
            lines.append(f"{phase}:")
            lines += [f"  {a.title} ({a.owner_role})" for a in actions]
        _bullets(slide, lines, top=top, font_size=13)

    # --- Reviewer notes
    if record.review:
        slide = new_slide()
        top = _title(slide, "Reviewer Notes")
        lines = [f"Status: {'PASSED' if record.review.passed else 'ISSUES FOUND'}"]
        if record.review.citation_coverage_pct is not None:
            lines.append(f"Citation coverage: {record.review.citation_coverage_pct:.1f}%")
        _bullets(slide, lines, top=top, font_size=16, height=Inches(1.1))
        if record.review.issues:
            rows = [[i.severity.value, i.stage, i.item_title, i.problem] for i in record.review.issues]
            _table(
                slide,
                ["Severity", "Stage", "Item", "Problem"],
                _truncate_rows(rows, 8),
                top=top + Inches(1.3),
                col_widths=[1.2, 2, 2.5, 4],
            )

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
